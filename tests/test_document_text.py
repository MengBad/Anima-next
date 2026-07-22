from __future__ import annotations

from io import BytesIO
import zipfile

import pytest

from anima.document_text import extract_document_text


def test_extracts_docx_paragraphs_and_tables():
    from docx import Document
    document = Document()
    document.add_paragraph("文档标题")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "列一"
    table.cell(0, 1).text = "列二"
    data = BytesIO()
    document.save(data)
    text = extract_document_text("notes.docx", data.getvalue())
    assert "文档标题" in text
    assert "列一 | 列二" in text


def test_extracts_xlsx_cells():
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "预算"
    sheet.append(["项目", "金额"])
    sheet.append(["服务器", 20])
    data = BytesIO()
    workbook.save(data)
    text = extract_document_text("budget.xlsx", data.getvalue())
    assert "[工作表: 预算]" in text
    assert "服务器 | 20" in text


def test_rejects_office_zip_bomb_before_parser():
    data = BytesIO()
    with zipfile.ZipFile(data, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (13 * 1024 * 1024))
    with pytest.raises(ValueError, match="解压"):
        extract_document_text("large.docx", data.getvalue())


def test_invalid_pdf_is_not_silently_treated_as_text():
    with pytest.raises(Exception):
        extract_document_text("bad.pdf", b"not a pdf")
