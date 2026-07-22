"""受限本地文档文本提取：PDF、DOCX、XLSX，不执行宏或外部转换器。"""
from __future__ import annotations

from io import BytesIO
import zipfile

MAX_DOCUMENT_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 12 * 1024 * 1024
MAX_TEXT_CHARS = 12_000
MAX_PDF_PAGES = 20
MAX_SHEET_CELLS = 5_000


def extract_document_text(name: str, content: bytes) -> str:
    """返回截断文本；调用者负责将失败信息转换为面向用户的降级说明。"""
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("附件超过 2MB 限制")
    suffix = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if suffix == "pdf":
        return _pdf(content)
    if suffix == "docx":
        _check_zip(content)
        return _docx(content)
    if suffix == "xlsx":
        _check_zip(content)
        return _xlsx(content)
    raise ValueError("不支持的文档格式")


def _pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(content))
    parts = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        parts.append(page.extract_text() or "")
        if sum(map(len, parts)) >= MAX_TEXT_CHARS:
            break
    return _limit("\n".join(parts))


def _docx(content: bytes) -> str:
    from docx import Document
    document = Document(BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            if sum(map(len, parts)) >= MAX_TEXT_CHARS:
                return _limit("\n".join(parts))
    return _limit("\n".join(parts))


def _xlsx(content: bytes) -> str:
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        parts: list[str] = []
        cells = 0
        for sheet in workbook.worksheets:
            parts.append(f"[工作表: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None and str(value).strip()]
                cells += len(values)
                if values:
                    parts.append(" | ".join(values))
                if cells >= MAX_SHEET_CELLS or sum(map(len, parts)) >= MAX_TEXT_CHARS:
                    return _limit("\n".join(parts))
        return _limit("\n".join(parts))
    finally:
        workbook.close()


def _check_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 2_000 or sum(info.file_size for info in infos) > MAX_UNZIPPED_BYTES:
                raise ValueError("Office 文档解压后超过安全限制")
    except zipfile.BadZipFile as exc:
        raise ValueError("不是有效的 Office 文档") from exc


def _limit(text: str) -> str:
    text = text.strip()
    return text[:MAX_TEXT_CHARS] + ("\n[内容已截断]" if len(text) > MAX_TEXT_CHARS else "")
